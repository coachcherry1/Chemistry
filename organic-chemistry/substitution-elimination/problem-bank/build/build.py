"""Write problems.json/csv and the student + answer-key Word documents."""
import os, sys, csv, json, warnings
warnings.filterwarnings('ignore')
HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.abspath(os.path.join(HERE, '..'))
sys.path.insert(0, HERE)
import problems as pr
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

P = pr.P
json.dump(P, open(os.path.join(ROOT, 'problems.json'), 'w'), indent=1, ensure_ascii=False)
with open(os.path.join(ROOT, 'problems.csv'), 'w', newline='', encoding='utf-8') as f:
    cols = ['id', 'type', 'mechanism', 'difficulty', 'question', 'answer', 'explanation', 'source', 'question_image', 'answer_image']
    w = csv.DictWriter(f, cols, extrasaction='ignore'); w.writeheader(); w.writerows(P)


def width(rel):
    px = Image.open(os.path.join(ROOT, rel)).size[0] / 2.0
    return min(6.0, px / 120.0)


def setup(doc, footer):
    sec = doc.sections[0]; sec.left_margin = sec.right_margin = Inches(0.9); sec.top_margin = sec.bottom_margin = Inches(0.8)
    st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(11)
    fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.add_run(footer + '  –  Page ').font.size = Pt(9)
    for tag, text in (('begin', None), (None, 'PAGE'), ('end', None)):
        run = fp.add_run(); run.font.size = Pt(9)
        el = OxmlElement('w:fldChar' if tag else 'w:instrText')
        if tag: el.set(qn('w:fldCharType'), tag)
        else: el.set(qn('xml:space'), 'preserve'); el.text = text
        run._r.append(el)


def block(doc):
    t = doc.add_table(rows=1, cols=1)
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    return t.cell(0, 0)


def student():
    d = Document(); setup(d, 'Substitution and Elimination Problems')
    h = d.add_paragraph(); h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run('Substitution and Elimination (SN1, SN2, E1, E2): 100 Problems'); r.bold = True; r.font.size = Pt(16)
    d.add_paragraph('Free response. For each problem, show structures clearly. Where asked, justify your answer with the four factors: '
                    'substrate, nucleophile/base, leaving group and solvent.')
    for p in P:
        c = block(d)
        para = c.paragraphs[0]; para.paragraph_format.space_before = Pt(8)
        a = para.add_run(f"{int(p['id'][1:])}.  "); a.bold = True
        b = para.add_run(f"[{p['type']}]  "); b.italic = True; b.font.size = Pt(9); b.font.color.rgb = RGBColor(0x55, 0x66, 0x77)
        para.add_run(p['question'])
        if p['question_image']:
            ip = c.add_paragraph(); ip.paragraph_format.left_indent = Inches(0.3)
            ip.add_run().add_picture(os.path.join(ROOT, p['question_image']), width=Inches(width(p['question_image'])))
        for _ in range(3):
            c.add_paragraph('_' * 86).paragraph_format.space_after = Pt(2)
    d.save(os.path.join(ROOT, 'SN_E_Problems.docx'))


def key():
    d = Document(); setup(d, 'Substitution and Elimination: Answer Key')
    h = d.add_paragraph(); h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run('Substitution and Elimination: Answer Key'); r.bold = True; r.font.size = Pt(16)
    n = d.add_paragraph('Substitution reasoning follows the four-factor framework in Klein, Organic Chemistry, Ch. 7. '
                        'Elimination items (E1/E2, Zaitsev/Hofmann, alcohol dehydration) use standard textbook conventions and were not '
                        'checked against an elimination chapter; compare them with your text. Items are tagged in the source field of problems.json.')
    n.runs[0].italic = True
    for p in P:
        c = block(d)
        para = c.paragraphs[0]; para.paragraph_format.space_before = Pt(8)
        a = para.add_run(f"{int(p['id'][1:])}.  {p['type']}  "); a.bold = True
        b = para.add_run(f"({p['mechanism']}; difficulty {p['difficulty']})"); b.italic = True; b.font.size = Pt(9)
        qp = c.add_paragraph(); qp.add_run('Q: ').bold = True; qp.add_run(p['question'])
        ap = c.add_paragraph(); r = ap.add_run('Answer: '); r.bold = True; r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79); ap.add_run(p['answer'])
        img = p['answer_image'] or p['question_image']
        if img:
            ip = c.add_paragraph(); ip.paragraph_format.left_indent = Inches(0.3)
            ip.add_run().add_picture(os.path.join(ROOT, img), width=Inches(width(img)))
        if p['explanation']:
            ep = c.add_paragraph(); ep.add_run('Why: ').bold = True; ep.add_run(p['explanation'])
        if p['source'] != pr.SRC_SN:
            sp = c.add_paragraph(); r = sp.add_run('Source note: ' + p['source']); r.italic = True; r.font.size = Pt(9)
    d.save(os.path.join(ROOT, 'SN_E_Answer_Key.docx'))


student(); key(); print('built', len(P), 'problems')
