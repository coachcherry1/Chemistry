"""One Word document per organic chapter (Ch. 2 to 11): a question document and a separate answer key.

Everything in the index for a chapter goes in (new question bank, acid/base quiz, substitution/elimination problems), grouped by textbook
section, multiple choice first. Written to organic-chemistry/word-docs/. Run from anywhere:  python3 tools/build_word_docs.py
"""
import os, re, sys, urllib.parse
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

ROOT = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), '..'))
sys.path.insert(0, os.path.join(ROOT, 'tools'))
import build_index as bi

OUT = os.path.join(ROOT, 'organic-chemistry', 'word-docs')
LEVEL = {1: 'Easy', 2: 'Basic', 3: 'Moderate', 4: 'Hard', 5: 'Very hard'}
PX_PER_INCH = 150


def slug(s):
    return re.sub(r'[^A-Za-z0-9]+', '_', s).strip('_')


def pic(doc_or_cell, rel, max_w):
    path = os.path.join(ROOT, urllib.parse.unquote(rel))
    if not os.path.exists(path):
        return False
    w = Image.open(path).size[0]
    width = min(max_w, w / PX_PER_INCH)
    para = doc_or_cell.add_paragraph()
    para.add_run().add_picture(path, width=Inches(width))
    return True


def new_doc(title, subtitle):
    d = Document()
    st = d.styles['Normal']
    st.font.name = 'Calibri'
    st.font.size = Pt(11)
    for s in d.sections:
        s.left_margin = s.right_margin = Inches(0.9)
        s.top_margin = s.bottom_margin = Inches(0.8)
    h = d.add_heading(title, level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p = d.add_paragraph(subtitle)
    p.runs[0].italic = True
    d.add_paragraph('Name: ______________________________    Date: ______________')
    return d


def section_heading(d, text):
    h = d.add_heading(text, level=2)
    h.paragraph_format.space_before = Pt(14)


def stem_para(d, n, it, show_id=True):
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f'{n}. ')
    r.bold = True
    stem = it['b'] or it['n']
    p.add_run(stem)
    tag = p.add_run(f'   [{it["id"]}, {LEVEL.get(int(it["dl"]) or 3, "")}]')
    tag.font.size = Pt(8)
    tag.font.color.rgb = RGBColor(0x80, 0x80, 0x80)


def options(d, it):
    imgs = any(o['t'] == 'img' for o in it['o'])
    if imgs:
        t = d.add_table(rows=2, cols=2)
        for k, o in enumerate(it['o']):
            cell = t.cell(k // 2, k % 2)
            pic(cell, o['v'], 2.8)      # the picture already has its letter drawn in
    else:
        for k, o in enumerate(it['o']):
            p = d.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.35)
            p.paragraph_format.space_after = Pt(0)
            p.add_run('ABCD'[k] + ')  ').bold = True
            p.add_run(o['v'])


def chapter_items(data, cid):
    secs = [s for t in bi.TAX if t['id'] == cid for s in t['sections']]
    order = {s['id']: i for i, s in enumerate(secs)}
    out = [d for d in data if d['c'] == 'organic' and d['k'] == 'question' and d['s'] and d['s'][0].split('.')[0] == cid]
    out.sort(key=lambda d: (order.get(d['s'][0], 999), d['f'] != 'Multiple choice', d['id']))
    return out, bi.SEC_LABEL


def build_chapter(data, tax):
    items, label = chapter_items(data, tax['id'])
    title = f"Chapter {tax['id']}: {tax['short']}"
    base = f"Ch{int(tax['id']):02d}_{slug(tax['short'])}"
    qd = new_doc(title, 'Organic chemistry questions. Multiple choice first within each textbook section.')
    kd = new_doc(title + ': Answer Key', 'Answers and explanations.')
    last = None
    for n, it in enumerate(items, 1):
        sec = it['s'][0]
        if sec != last:
            section_heading(qd, label[sec])
            section_heading(kd, label[sec])
            last = sec
        stem_para(qd, n, it)
        stem_para(kd, n, it)
        if it['im']:
            pic(qd, it['im'], 5.4)
        if it['f'] == 'Multiple choice':
            options(qd, it)
            k = kd.add_paragraph()
            k.add_run('Answer: ').bold = True
            idx = 'ABCD'.index(it['oa'])
            o = it['o'][idx]
            k.add_run(f"{it['oa']}" + (f") {o['v']}" if o['t'] == 'text' else ''))
            if o['t'] == 'img':
                pic(kd, o['v'], 2.8)
        else:
            qd.add_paragraph('\n\n')
            k = kd.add_paragraph()
            k.add_run('Answer: ').bold = True
            k.add_run(it['a'])
            if it['am']:
                pic(kd, it['am'], 5.0)
        if it['w']:
            e = kd.add_paragraph()
            r = e.add_run('Why: ')
            r.italic = True
            e.add_run(it['w'])
    qn, kn = base + '_Questions.docx', base + '_Answer_Key.docx'
    qd.save(os.path.join(OUT, qn))
    kd.save(os.path.join(OUT, kn))
    return len(items), qn, kn


def main():
    os.makedirs(OUT, exist_ok=True)
    data = [it for loader in bi.LOADERS for it in loader()]
    rows = []
    for t in bi.TAX:
        if t['id'] == 'X':
            continue
        n, qn, kn = build_chapter(data, t)
        rows.append((t['id'], t['short'], n, qn, kn))
        print(f"Ch {t['id']}: {n} questions -> {qn}")
    lines = ['# Organic chemistry Word documents', '',
             'One question document and one answer key per chapter (Klein Ch. 2 to 11). Download from the table, or from the **Downloads** button on the index page.', '',
             '| Chapter | Questions | Question document | Answer key |', '| --- | --- | --- | --- |']
    for cid, short, n, qn, kn in rows:
        lines.append(f'| Ch. {cid} {short} | {n} | [{qn}]({qn}) | [{kn}]({kn}) |')
    lines += ['', 'Made by `python3 tools/build_word_docs.py` (it reads the same data as the index, so run it after any wording or difficulty change).', '']
    open(os.path.join(OUT, 'README.md'), 'w', encoding='utf-8').write('\n'.join(lines))


if __name__ == '__main__':
    main()
